import docx, random
import GUI

doc = docx.Document('0.docx')

table = doc.add_table(rows = 3, cols = 3, style='Table Grid')

for row in range(3):
    for col in range(3):
        cell = table.cell(row, col)
        cell.text = str(random.randint(0, 100))

dir = GUI.lbl['text']
print(dir)
doc.save(dir + '/' + '0.docx')